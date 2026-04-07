from __future__ import annotations

import csv
import json
from collections.abc import Iterable, Iterator
from pathlib import Path
from typing import Any

from sqlalchemy.orm import Session

from app.core.settings import get_settings
from app.services.query_service import QueryService
from app.services.solution_repository import SolutionRepositoryService

SKIP_PARAMETER_PLOTS = {"imaging_time"}


class ExportService:
    def __init__(self, db: Session):
        self.db = db
        self.settings = get_settings()
        self.query = QueryService(db)

    def export_solution_repository(self, export_format: str = "json") -> str:
        return SolutionRepositoryService(self.db).export_records(export_format=export_format)

    def export_events_csv(
        self,
        task_id: int,
        task_uuid: str,
        *,
        side_scopes: list[str] | None = None,
        side_groups: list[str] | None = None,
        chip_names: list[str] | None = None,
    ) -> str:
        output = self._export_path(f"{task_uuid}_events.csv")
        self._write_csv_stream(
            output,
            self._iter_event_rows(
                task_id,
                side_scopes=side_scopes,
                side_groups=side_groups,
                chip_names=chip_names,
            ),
        )
        return str(output)

    def export_error_report_csv(
        self,
        task_id: int,
        task_uuid: str,
        *,
        side_scopes: list[str] | None = None,
        side_groups: list[str] | None = None,
        chip_names: list[str] | None = None,
    ) -> str:
        rows = self.query.get_error_clusters(
            task_id=task_id,
            limit=100000,
            offset=0,
            side_scopes=side_scopes,
            side_groups=side_groups,
            chip_names=chip_names,
        )["items"]
        output = self._export_path(f"{task_uuid}_errors.csv")
        self._write_csv(output, rows)
        return str(output)

    def export_parameter_results_csv(
        self,
        task_id: int,
        task_uuid: str,
        *,
        side_scopes: list[str] | None = None,
        side_groups: list[str] | None = None,
        chip_names: list[str] | None = None,
    ) -> str:
        rows = [
            row
            for row in self.query.get_parameter_results(task_id)
            if self.query._row_matches_scope_filters(
                row,
                side_scopes=side_scopes,
                side_groups=side_groups,
                chip_names=chip_names,
            )
        ]
        output = self._export_path(f"{task_uuid}_parameter_results.csv")
        self._write_csv(output, rows)
        return str(output)

    def export_parameter_series_csv(
        self,
        task_id: int,
        task_uuid: str,
        parameter_name: str,
        unit: str = "s",
        *,
        side_scopes: list[str] | None = None,
        side_groups: list[str] | None = None,
        chip_names: list[str] | None = None,
    ) -> str:
        rows = self.query.get_parameter_series(
            task_id,
            parameter_name,
            unit=unit,
            side_scopes=side_scopes,
            side_groups=side_groups,
            chip_names=chip_names,
        )
        output = self._export_path(f"{task_uuid}_trend_{parameter_name}_{unit}.csv")
        self._write_csv(output, rows)
        return str(output)

    def export_row_scan_metric_series_csv(
        self,
        task_id: int,
        task_uuid: str,
        unit: str = "ms",
        *,
        side_scopes: list[str] | None = None,
        side_groups: list[str] | None = None,
        chip_names: list[str] | None = None,
    ) -> str:
        rows = self.query.get_row_scan_metric_stage_series(
            task_id,
            unit=unit,
            side_scopes=side_scopes,
            side_groups=side_groups,
            chip_names=chip_names,
        )
        output = self._export_path(f"{task_uuid}_trend_row_scan_metric_{unit}.csv")
        self._write_csv(output, rows)
        return str(output)

    def export_json_report(
        self,
        task_id: int,
        task_uuid: str,
        *,
        side_scopes: list[str] | None = None,
        side_groups: list[str] | None = None,
        chip_names: list[str] | None = None,
    ) -> str:
        payload = self._build_report_payload(
            task_id,
            side_scopes=side_scopes,
            side_groups=side_groups,
            chip_names=chip_names,
        )
        output = self._export_path(f"{task_uuid}_report.json")
        output.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        return str(output)

    def export_html_report(
        self,
        task_id: int,
        task_uuid: str,
        *,
        side_scopes: list[str] | None = None,
        side_groups: list[str] | None = None,
        chip_names: list[str] | None = None,
    ) -> str:
        payload = self._build_report_payload(
            task_id,
            side_scopes=side_scopes,
            side_groups=side_groups,
            chip_names=chip_names,
        )
        output = self._export_path(f"{task_uuid}_report.html")
        output.write_text(self._build_html(payload, task_uuid), encoding="utf-8")
        return str(output)

    def export_excel_report(
        self,
        task_id: int,
        task_uuid: str,
        *,
        side_scopes: list[str] | None = None,
        side_groups: list[str] | None = None,
        chip_names: list[str] | None = None,
    ) -> str:
        try:
            from openpyxl import Workbook
            from openpyxl.utils import get_column_letter
        except ModuleNotFoundError as exc:
            raise RuntimeError("Excel export dependency missing: install openpyxl") from exc

        payload = self._build_report_payload(
            task_id,
            side_scopes=side_scopes,
            side_groups=side_groups,
            chip_names=chip_names,
        )
        output = self._export_path(f"{task_uuid}_report.xlsx")
        wb = Workbook()
        ws = wb.active
        if ws is None:
            raise RuntimeError("Excel workbook initialization failed")
        ws.title = "Dashboard"
        ws.append(["metric", "value"])
        for key, value in payload["dashboard"].items():
            if isinstance(value, list):
                continue
            ws.append([key, self._safe_value(value)])
        self._autosize_sheet(ws, get_column_letter)
        self._add_sheet(wb, "ScopeCatalog", payload.get("scope_catalog", {}), get_column_letter)
        self._add_sheet(wb, "Events", payload["events"][:5000], get_column_letter)
        self._add_sheet(wb, "Errors", payload["errors"], get_column_letter)
        self._add_sheet(wb, "CycleSummary", payload["cycle_summary"], get_column_letter)
        self._add_sheet(wb, "Timeline", payload["timeline"][:5000], get_column_letter)
        self._add_sheet(wb, "SubstepCycle", payload["substep_cycle"], get_column_letter)
        self._add_sheet(wb, "RowScanMetric", payload["row_scan_metric_series"], get_column_letter)
        self._add_sheet(wb, "Audit", payload["audit_logs"], get_column_letter)
        self._add_sheet(wb, "LLM", payload["llm_results"], get_column_letter)
        for name, rows in payload["parameter_series"].items():
            self._add_sheet(wb, f"trend_{name}"[:31], rows, get_column_letter)
        wb.save(output)
        return str(output)

    def export_pdf_report(
        self,
        task_id: int,
        task_uuid: str,
        *,
        side_scopes: list[str] | None = None,
        side_groups: list[str] | None = None,
        chip_names: list[str] | None = None,
    ) -> str:
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.utils import simpleSplit
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.pdfgen import canvas
        except ModuleNotFoundError as exc:
            raise RuntimeError("PDF export dependency missing: install reportlab") from exc

        payload = self._build_report_payload(
            task_id,
            side_scopes=side_scopes,
            side_groups=side_groups,
            chip_names=chip_names,
        )
        output = self._export_path(f"{task_uuid}_report.pdf")
        c = canvas.Canvas(str(output), pagesize=A4)
        width, height = A4
        margin_x = 40
        y = height - 40
        title_font, body_font = self._register_pdf_fonts(pdfmetrics, UnicodeCIDFont)
        c.setTitle(f"Sequencer Log Report - {task_uuid}")
        y = self._draw_wrapped_line(c, f"Sequencer Log Report - {task_uuid}", margin_x, y, width - 2 * margin_x, simpleSplit, title_font, 14, 18)
        dash = payload["dashboard"]
        lines = [
            "",
            "Summary",
            f"Files: {dash.get('file_count', 0)}",
            f"Events: {dash.get('total_events', 0)}",
            f"Errors: {dash.get('total_errors', 0)}",
            f"Unique errors: {dash.get('unique_error_count', 0)}",
        ]
        for line in lines:
            y = self._draw_wrapped_line(c, line, margin_x, y, width - 2 * margin_x, simpleSplit, body_font, 10, 14)
        c.save()
        return str(output)

    def _build_report_payload(
        self,
        task_id: int,
        *,
        side_scopes: list[str] | None = None,
        side_groups: list[str] | None = None,
        chip_names: list[str] | None = None,
    ) -> dict[str, Any]:
        defs = [d for d in self.query.get_parameter_definitions() if d["parameter_name"] not in SKIP_PARAMETER_PLOTS]
        payload: dict[str, Any] = {}
        jobs = {
            "dashboard": ("get_dashboard", {}),
            "events": (
                "list_events",
                {
                    "limit": 2000,
                    "offset": 0,
                    "side_scopes": side_scopes,
                    "side_groups": side_groups,
                    "chip_name": chip_names,
                },
            ),
            "errors": (
                "get_error_clusters",
                {
                    "limit": 500,
                    "offset": 0,
                    "side_scopes": side_scopes,
                    "side_groups": side_groups,
                    "chip_names": chip_names,
                },
            ),
            "cycle_summary": (
                "get_cycle_summaries",
                {
                    "unit": "s",
                    "side_scopes": side_scopes,
                    "side_groups": side_groups,
                    "chip_names": chip_names,
                },
            ),
            "timeline": (
                "get_movement_timeline",
                {
                    "track_order": "cycle",
                    "track_granularity": "side_chip",
                    "side_scopes": side_scopes,
                    "side_groups": side_groups,
                    "chip_names": chip_names,
                },
            ),
            "substep_cycle": (
                "get_substep_cycle_series",
                {
                    "agg_mode": "mean",
                    "unit": "s",
                    "side_scopes": side_scopes,
                    "side_groups": side_groups,
                    "chip_names": chip_names,
                },
            ),
            "row_scan_metric_series": (
                "get_row_scan_metric_stage_series",
                {
                    "unit": "s",
                    "side_scopes": side_scopes,
                    "side_groups": side_groups,
                    "chip_names": chip_names,
                },
            ),
            "audit_logs": ("get_audit_logs", {}),
            "llm_results": ("get_llm_results", {}),
        }
        for key, (fn, kwargs) in jobs.items():
            value = getattr(self.query, fn)(task_id, **kwargs)
            payload[key] = value["items"] if key in {"events", "errors"} else value

        param_series: dict[str, list[dict[str, Any]]] = {}
        for definition in defs:
            param_series[definition["parameter_name"]] = self.query.get_parameter_series(
                task_id,
                definition["parameter_name"],
                unit="s",
                side_scopes=side_scopes,
                side_groups=side_groups,
                chip_names=chip_names,
            )

        payload["parameter_definitions"] = defs
        payload["parameter_series"] = param_series
        payload["scope_catalog"] = self.query.get_scope_catalog(task_id)
        return payload

    def _iter_event_rows(
        self,
        task_id: int,
        page_size: int = 2000,
        *,
        side_scopes: list[str] | None = None,
        side_groups: list[str] | None = None,
        chip_names: list[str] | None = None,
    ) -> Iterator[dict[str, Any]]:
        offset = 0
        while True:
            payload = self.query.list_events(
                task_id=task_id,
                limit=page_size,
                offset=offset,
                side_scopes=side_scopes,
                side_groups=side_groups,
                chip_name=chip_names,
            )
            rows = payload.get("items") or []
            if not rows:
                break
            yield from rows
            if len(rows) < page_size:
                break
            offset += len(rows)

    @staticmethod
    def _plotly():
        import plotly.express as px
        import plotly.io as pio

        return px, pio

    def _build_html(self, payload: dict[str, Any], task_uuid: str) -> str:
        def section(title: str, content: str, sid: str) -> str:
            return f'<div class="section"><h2 id="{sid}">{title}</h2>{content}</div>'

        param_html = ""
        for name, rows in sorted(payload["parameter_series"].items()):
            param_html += self._plot_parameter_series(name, rows)
            param_html += self._table(rows[:100], f"{name} Trend Sample")
        param_html += self._plot_metric_series(payload["row_scan_metric_series"])
        param_html += self._table(payload["row_scan_metric_series"][:200], "Row Scan Metric Summary")

        sections = [
            section(
                "Overview",
                f"<h1>Sequencer Unified HTML Report</h1><p><strong>Task UUID:</strong> {task_uuid}</p>"
                "<p>This report includes events, timing, parameter trends, errors, timeline, and substep-cycle views.</p>",
                "home",
            ),
            section("Scope Catalog", self._table(self._scope_catalog_rows(payload.get("scope_catalog", {})), "Scope Catalog"), "scope"),
            section("Events", self._table(payload["events"][:200], "Event Summary (Top 200)"), "events"),
            section("Timing", self._plot_cycle_summary(payload["cycle_summary"]) + self._table(payload["cycle_summary"], "Cycle Summary"), "timing"),
            section("Parameters", param_html, "params"),
            section("Errors", self._plot_errors(payload["errors"]) + self._table(payload["errors"][:50], "Error Top 50"), "errors"),
            section("Timeline", self._plot_timeline(payload["timeline"]) + self._table(payload["timeline"][:150], "Timeline Summary"), "timeline"),
            section("Substep-Cycle", self._plot_substep_cycle(payload["substep_cycle"]) + self._table(payload["substep_cycle"][:200], "Substep-Cycle Summary"), "substep"),
        ]
        toc = (
            '<div class="section"><h2>Contents</h2><ul>'
            '<li><a href="#scope">Scope Catalog</a></li>'
            '<li><a href="#events">Events</a></li>'
            '<li><a href="#timing">Timing</a></li>'
            '<li><a href="#params">Parameters</a></li>'
            '<li><a href="#errors">Errors</a></li>'
            '<li><a href="#timeline">Timeline</a></li>'
            '<li><a href="#substep">Substep-Cycle</a></li>'
            "</ul></div>"
        )
        style = (
            '<style>body{font-family:Arial,"Microsoft YaHei",sans-serif;margin:18px;background:#f7f9fc;color:#111827}'
            '.section{background:white;padding:16px 20px;border-radius:10px;box-shadow:0 1px 4px rgba(0,0,0,.08);margin-bottom:16px}'
            'table{border-collapse:collapse;width:100%;font-size:12px}th,td{border:1px solid #dbe3ee;padding:6px 8px;vertical-align:top}'
            'th{background:#eff6ff}h1,h2{margin:0 0 12px 0}a{color:#2563eb}</style>'
        )
        body = toc + "".join(sections)
        return f'<!DOCTYPE html><html><head><meta charset="utf-8"/><title>Sequencer Report</title>{style}</head><body>{body}</body></html>'

    def _plot_cycle_summary(self, rows: list[dict[str, Any]]) -> str:
        if not rows:
            return "<p>No cycle summary.</p>"
        px, pio = self._plotly()
        color_field = "chip_name" if any(row.get("chip_name") for row in rows) else ("side_scope" if any(row.get("side_scope") for row in rows) else None)
        kwargs: dict[str, Any] = {"x": "cycle_no", "y": "total_duration_value", "markers": True, "title": "Cycle Duration Trend"}
        if color_field:
            kwargs["color"] = color_field
        fig = px.line(rows, **kwargs)
        return pio.to_html(fig, include_plotlyjs="inline", full_html=False)

    def _plot_parameter_series(self, name: str, rows: list[dict[str, Any]]) -> str:
        if name in SKIP_PARAMETER_PLOTS:
            return ""
        if not rows:
            return f"<p>{name}: no data.</p>"
        px, pio = self._plotly()
        x_col = "cycle"
        plot_rows = rows
        if rows[0].get("x_axis_type") == "time":
            x_col = "start_time"
            plot_rows = sorted(rows, key=lambda r: (r.get("start_time") or "", r.get("cycle") if r.get("cycle") is not None else -1))
        color_field = "chip_name" if any(row.get("chip_name") for row in rows) else ("side_scope" if any(row.get("side_scope") for row in rows) else None)
        kwargs: dict[str, Any] = {"x": x_col, "y": "duration_value", "markers": True, "title": f"{name} Trend"}
        if color_field:
            kwargs["color"] = color_field
        fig = px.line(plot_rows, **kwargs)
        threshold_value = next((row.get("threshold_value") for row in plot_rows if row.get("threshold_value") is not None), None)
        expected_value = next((row.get("expected_value") for row in plot_rows if row.get("expected_value") is not None), None)
        if threshold_value is not None:
            fig.add_hline(y=threshold_value, line_dash="dash", line_color="red")
        if expected_value is not None:
            fig.add_hline(y=expected_value, line_dash="dot", line_color="green")
        return pio.to_html(fig, include_plotlyjs=False, full_html=False)

    def _plot_metric_series(self, rows: list[dict[str, Any]]) -> str:
        if not rows:
            return "<p>No row scan metrics.</p>"
        px, pio = self._plotly()
        color_field = "metric_stage"
        if any(row.get("chip_name") for row in rows):
            for row in rows:
                row.setdefault("_metric_scope", f"{row.get('metric_stage') or 'metric'} | {row.get('chip_name') or row.get('side_scope') or 'Unassigned'}")
            color_field = "_metric_scope"
        fig = px.line(rows, x="cycle", y="duration_value", color=color_field, markers=True, title="Row Scan Metric Trend")
        return pio.to_html(fig, include_plotlyjs=False, full_html=False)

    def _plot_errors(self, rows: list[dict[str, Any]]) -> str:
        if not rows:
            return "<p>No error data.</p>"
        px, pio = self._plotly()
        color_field = "primary_side_scope" if any(row.get("primary_side_scope") for row in rows) else ("error_family_display" if "error_family_display" in rows[0] else "error_family")
        fig = px.bar(rows[:20], x="display_signature", y="count", color=color_field, title="Error Top 20")
        fig.update_xaxes(tickangle=-20)
        return pio.to_html(fig, include_plotlyjs=False, full_html=False)

    def _plot_timeline(self, rows: list[dict[str, Any]]) -> str:
        if not rows:
            return "<p>No timeline data.</p>"
        px, pio = self._plotly()
        color_field = "side_scope" if any(row.get("side_scope") for row in rows) else "sub_step"
        fig = px.timeline(
            rows,
            x_start="start",
            x_end="end",
            y="track",
            color=color_field,
            hover_data=["sub_step", "side_scope", "chip_name", "chuck_no", "cycle_no", "duration_ms", "component", "message"],
        )
        fig.update_layout(height=min(max(500, 22 * len({row["track"] for row in rows}) + 180), 2400))
        return pio.to_html(fig, include_plotlyjs=False, full_html=False)

    def _plot_substep_cycle(self, rows: list[dict[str, Any]]) -> str:
        if not rows:
            return "<p>No substep-cycle data.</p>"
        px, pio = self._plotly()
        aggregates: dict[str, float] = {}
        counts: dict[str, int] = {}
        for row in rows:
            sub_step = str(row.get("sub_step") or "Unknown")
            aggregates[sub_step] = aggregates.get(sub_step, 0.0) + float(row.get("duration_value") or 0.0)
            counts[sub_step] = counts.get(sub_step, 0) + 1
        top_substeps = sorted(aggregates.keys(), key=lambda key: aggregates[key] / max(counts[key], 1), reverse=True)[:12]
        plot_rows = [row for row in rows if str(row.get("sub_step") or "Unknown") in top_substeps]
        color_field = "sub_step"
        if any(row.get("chip_name") for row in plot_rows):
            for row in plot_rows:
                row.setdefault("_substep_scope", f"{row.get('sub_step') or 'Unknown'} | {row.get('chip_name') or row.get('side_scope') or 'Unassigned'}")
            color_field = "_substep_scope"
        fig = px.line(plot_rows, x="cycle_no", y="duration_value", color=color_field, title="Substep-Cycle Trend (Top 12 by Avg Duration)")
        return pio.to_html(fig, include_plotlyjs=False, full_html=False)

    def _table(self, rows: list[dict[str, Any]], title: str) -> str:
        if not rows:
            return f"<h3>{title}</h3><p>No data.</p>"
        headers = list(rows[0].keys())
        head = "".join(f"<th>{header}</th>" for header in headers)
        body = []
        for row in rows:
            body.append("<tr>" + "".join(f"<td>{self._safe_value(row.get(header))}</td>" for header in headers) + "</tr>")
        return f'<h3>{title}</h3><table><thead><tr>{head}</tr></thead><tbody>{"".join(body)}</tbody></table>'

    def _scope_catalog_rows(self, catalog: dict[str, Any]) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        for side in catalog.get("sides", []) or []:
            rows.append({"scope_type": "side", **dict(side)})
        for chip in catalog.get("chips", []) or []:
            rows.append({"scope_type": "chip", **dict(chip)})
        return rows

    def _export_path(self, filename: str) -> Path:
        output = Path(self.settings.export_dir) / filename
        output.parent.mkdir(parents=True, exist_ok=True)
        return output

    @staticmethod
    def _safe_value(value: Any) -> Any:
        if value is None:
            return ""
        if isinstance(value, (str, int, float, bool)):
            return value
        if isinstance(value, list):
            if value and all(isinstance(item, dict) for item in value):
                return json.dumps(value, ensure_ascii=False, default=str)
            return ", ".join(str(item) for item in value)
        if isinstance(value, dict):
            return json.dumps(value, ensure_ascii=False, default=str)
        return str(value)

    @classmethod
    def _write_csv(cls, output: Path, rows: list[dict[str, Any]]) -> None:
        if not rows:
            output.write_text("", encoding="utf-8")
            return
        with output.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            for row in rows:
                writer.writerow({key: cls._safe_value(value) for key, value in row.items()})

    @classmethod
    def _write_csv_stream(cls, output: Path, rows: Iterable[dict[str, Any]]) -> None:
        iterator = iter(rows)
        first = next(iterator, None)
        if first is None:
            output.write_text("", encoding="utf-8")
            return
        with output.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(first.keys()))
            writer.writeheader()
            writer.writerow({key: cls._safe_value(value) for key, value in first.items()})
            for row in iterator:
                writer.writerow({key: cls._safe_value(value) for key, value in row.items()})

    @classmethod
    def _add_sheet(cls, wb, name: str, rows: Any, get_column_letter) -> None:
        ws = wb.create_sheet(title=name[:31])
        if isinstance(rows, dict):
            rows = [rows]
        if not rows:
            ws.append(["empty"])
            return
        if isinstance(rows, list) and rows and isinstance(rows[0], dict):
            headers = list(rows[0].keys())
            ws.append(headers)
            for row in rows:
                ws.append([cls._safe_value(row.get(header)) for header in headers])
        else:
            ws.append(["value"])
            for row in rows if isinstance(rows, list) else [rows]:
                ws.append([cls._safe_value(row)])
        cls._autosize_sheet(ws, get_column_letter)

    @staticmethod
    def _autosize_sheet(ws, get_column_letter) -> None:
        max_widths: dict[int, int] = {}
        for row in ws.iter_rows():
            for cell in row:
                value = "" if cell.value is None else str(cell.value)
                max_widths[cell.column] = max(max_widths.get(cell.column, 0), min(len(value), 80))
        for idx, width in max_widths.items():
            ws.column_dimensions[get_column_letter(idx)].width = max(12, min(width + 2, 80))

    @staticmethod
    def _register_pdf_fonts(pdfmetrics, UnicodeCIDFont) -> tuple[str, str]:
        try:
            pdfmetrics.getFont("STSong-Light")
        except KeyError:
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        return "STSong-Light", "STSong-Light"

    @staticmethod
    def _draw_wrapped_line(c, text_line: str, x: float, y: float, max_width: float, simpleSplit, font_name: str, font_size: int, leading: int = 14) -> float:
        wrapped = simpleSplit(str(text_line or ""), font_name, font_size, max_width) or [""]
        c.setFont(font_name, font_size)
        for segment in wrapped:
            if y < 40:
                c.showPage()
                c.setFont(font_name, font_size)
                y = 800
            c.drawString(x, y, segment)
            y -= leading
        return y
