from __future__ import annotations

import csv
import json
import re
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Any

from sqlalchemy import or_, select, text
from sqlalchemy.orm import Session

from app.core.security import utcnow
from app.core.settings import get_settings
from app.models.db_models import (
    SolutionMessageKeywordModel,
    SolutionModuleLinkModel,
    SolutionRecordModel,
    SolutionTagLinkModel,
    SolutionTagModel,
    SolutionTaskClusterLinkModel,
    SolutionTaskLinkModel,
    TaskClusterModel,
)
from app.services.error_code_service import ErrorCodeService
from app.services.solution_catalog_service import SYSTEM_MODULES, SolutionCatalogService

MODULE_PREFIXES: dict[str, str] = {item["module_key"]: item["prefix"] for item in SYSTEM_MODULES}
MODULE_TREE: list[dict[str, Any]] = [{"name": item["module_key"], "children": [item["display_name"]]} for item in SYSTEM_MODULES]
TOKEN_RE = re.compile(r"[A-Za-z0-9_\u4e00-\u9fff]{2,}")


def _json_dump(data: Any) -> str | None:
    if data in (None, "", [], {}):
        return None
    return json.dumps(data, ensure_ascii=False)


def _json_load(text_value: str | None) -> Any:
    if not text_value:
        return None
    try:
        return json.loads(text_value)
    except Exception:
        return text_value


def _preview_text(value: str | None, limit: int = 160) -> str | None:
    if not value:
        return value
    text_value = " ".join(str(value).split())
    return text_value[:limit]


def _coerce_string_list(value: Any) -> list[str]:
    if value in (None, "", []):
        return []
    if isinstance(value, str):
        return [item.strip() for item in value.split(",") if item and item.strip()]
    if isinstance(value, (list, tuple, set)):
        return [str(item).strip() for item in value if str(item).strip()]
    return [str(value).strip()]


def _tokenize_keywords(*values: Any) -> list[str]:
    tokens: list[str] = []
    for value in values:
        for token in TOKEN_RE.findall(str(value or "")):
            lowered = token.strip().lower()
            if len(lowered) >= 2:
                tokens.append(lowered)
    return list(dict.fromkeys(tokens))


class SolutionRepositoryService:
    def __init__(self, db: Session):
        self.db = db
        self.settings = get_settings()
        self.catalog_service = SolutionCatalogService(db)
        self.error_code_service = ErrorCodeService(db)

    def module_tree(self) -> list[dict[str, Any]]:
        return MODULE_TREE

    def module_prefixes(self) -> dict[str, str]:
        return MODULE_PREFIXES

    def fts_enabled(self) -> bool:
        try:
            row = self.db.execute(
                text("SELECT name FROM sqlite_master WHERE type='table' AND name='solution_search_fts'")
            ).first()
            return bool(row)
        except Exception:
            return False

    def validate_case_payload(self, payload: dict[str, Any]) -> dict[str, Any]:
        data = dict(payload or {})
        module_value = str(data.get("module") or data.get("module_key") or data.get("module_prefix") or "").strip()
        if not module_value:
            raise ValueError("module 不能为空")
        module_row = self.catalog_service.resolve_module(module_value)

        error_name = str(data.get("error_name") or "").strip()
        if not error_name:
            raise ValueError("error_name 不能为空")

        normalized_signature = str(data.get("normalized_signature") or "").strip() or None
        message = str(data.get("message") or "").strip() or None
        if not normalized_signature and not message and not str(data.get("trigger_scenario") or "").strip():
            raise ValueError("message、normalized_signature、trigger_scenario 至少提供一个")

        reusable = data.get("reusable", True)
        if isinstance(reusable, str):
            reusable = reusable.strip().lower() not in {"false", "0", "no", "n"}

        existing_error_code = str(data.get("error_code") or "").strip() or None
        if existing_error_code:
            expected_prefix = module_row.prefix.upper()
            if not existing_error_code.upper().startswith(f"{expected_prefix}-"):
                raise ValueError(f"错误码 {existing_error_code} 与模块前缀 {expected_prefix} 不匹配")
            duplicate = self.db.scalar(
                select(SolutionRecordModel.id).where(
                    SolutionRecordModel.error_code == existing_error_code,
                    SolutionRecordModel.id != int(data.get("id") or 0),
                )
            )
            if duplicate:
                raise ValueError(f"错误码 {existing_error_code} 已存在")

        task_clusters = _coerce_string_list(data.get("task_clusters") or data.get("task_cluster"))
        tags = _coerce_string_list(data.get("tags"))
        additional_modules = _coerce_string_list(data.get("additional_modules"))
        message_keywords = _coerce_string_list(data.get("message_keywords"))
        if not message_keywords:
            message_keywords = _tokenize_keywords(message, error_name, data.get("trigger_scenario"))

        task_links = list(data.get("task_links") or [])
        if data.get("task_uuid") or normalized_signature:
            task_links.append(
                {
                    "task_uuid": str(data.get("task_uuid") or "").strip() or None,
                    "normalized_signature": normalized_signature,
                    "relation_type": "reference",
                }
            )

        return {
            "id": int(data.get("id") or 0) or None,
            "error_name": error_name,
            "error_category": str(data.get("error_category") or "").strip() or None,
            "module": module_row.module_key,
            "submodule": str(data.get("submodule") or "").strip() or None,
            "error_code": existing_error_code,
            "error_code_prefix": module_row.prefix,
            "message": message,
            "normalized_signature": normalized_signature,
            "exception_description": str(data.get("exception_description") or "").strip() or None,
            "trigger_scenario": str(data.get("trigger_scenario") or "").strip() or None,
            "impact_scope": str(data.get("impact_scope") or "").strip() or None,
            "report_source": str(data.get("report_source") or "").strip() or None,
            "related_logs": _json_dump(data.get("related_logs")),
            "related_source_files": _json_dump(data.get("related_source_files")),
            "root_cause_analysis": str(data.get("root_cause_analysis") or "").strip() or None,
            "verified_solution": str(data.get("verified_solution") or "").strip() or None,
            "workaround": str(data.get("workaround") or "").strip() or None,
            "owner_department": str(data.get("owner_department") or "").strip() or None,
            "submitter": str(data.get("submitter") or "").strip() or None,
            "source": str(data.get("source") or "").strip() or None,
            "review_status": str(data.get("review_status") or "approved").strip() or "approved",
            "reusable": bool(reusable),
            "similar_case_refs": _json_dump(data.get("similar_case_refs")),
            "metadata_json": _json_dump(data.get("metadata")),
            "task_clusters": task_clusters,
            "tags": tags,
            "message_keywords": message_keywords,
            "task_links": task_links,
            "additional_modules": additional_modules,
        }

    def create_record(self, payload: dict[str, Any], *, actor: str | None = None) -> dict[str, Any]:
        data = self.validate_case_payload(payload)
        now = utcnow()
        error_code = data["error_code"] or self.error_code_service.allocate(data["error_code_prefix"])
        row = SolutionRecordModel(
            error_name=data["error_name"],
            error_category=data["error_category"],
            module=data["module"],
            submodule=data["submodule"],
            error_code=error_code,
            error_code_prefix=data["error_code_prefix"],
            message=data["message"],
            normalized_signature=data["normalized_signature"],
            exception_description=data["exception_description"],
            trigger_scenario=data["trigger_scenario"],
            impact_scope=data["impact_scope"],
            report_source=data["report_source"],
            related_logs=data["related_logs"],
            related_source_files=data["related_source_files"],
            root_cause_analysis=data["root_cause_analysis"],
            verified_solution=data["verified_solution"],
            workaround=data["workaround"],
            owner_department=data["owner_department"],
            submitter=data["submitter"] or actor,
            source=data["source"],
            review_status=data["review_status"],
            reusable=data["reusable"],
            similar_case_refs=data["similar_case_refs"],
            metadata_json=data["metadata_json"],
            created_at=now,
            updated_at=now,
        )
        self.db.add(row)
        self.db.flush()
        self._sync_relations(row.id, data=data, actor=actor or data["submitter"] or "system")
        self._sync_search_index(row.id)
        self.db.commit()
        self.db.refresh(row)
        return self.get_record(row.id)

    def update_record(self, record_id: int, payload: dict[str, Any], *, actor: str | None = None) -> dict[str, Any]:
        row = self.db.get(SolutionRecordModel, record_id)
        if not row:
            raise KeyError("solution record not found")
        base_payload = {**self.serialize(row), **payload, "id": record_id, "error_code": payload.get("error_code", row.error_code)}
        data = self.validate_case_payload(base_payload)
        for field_name in [
            "error_name",
            "error_category",
            "module",
            "submodule",
            "error_code_prefix",
            "message",
            "normalized_signature",
            "exception_description",
            "trigger_scenario",
            "impact_scope",
            "report_source",
            "related_logs",
            "related_source_files",
            "root_cause_analysis",
            "verified_solution",
            "workaround",
            "owner_department",
            "submitter",
            "source",
            "review_status",
            "reusable",
            "similar_case_refs",
            "metadata_json",
        ]:
            setattr(row, field_name, data[field_name])
        row.updated_at = utcnow()
        self.db.flush()
        self._sync_relations(row.id, data=data, actor=actor or data["submitter"] or "system")
        self._sync_search_index(row.id)
        self.db.commit()
        self.db.refresh(row)
        return self.get_record(row.id)

    def get_record(self, record_id: int) -> dict[str, Any]:
        row = self.db.get(SolutionRecordModel, record_id)
        if not row:
            raise KeyError("solution record not found")
        relation_maps = self._load_relation_maps([row.id])
        return self.serialize(row, relation_maps=relation_maps)

    def list_records(
        self,
        *,
        module: str | None = None,
        submodule: str | None = None,
        error_code: str | None = None,
        error_name: str | None = None,
        message: str | None = None,
        message_keyword: str | None = None,
        normalized_signature: str | None = None,
        trigger_scenario: str | None = None,
        task_cluster: str | None = None,
        submitter: str | None = None,
        reusable: bool | None = None,
        review_status: str | None = None,
        search: str | None = None,
        created_from: str | None = None,
        created_to: str | None = None,
        updated_from: str | None = None,
        updated_to: str | None = None,
        viewer_username: str | None = None,
        viewer_is_reviewer: bool = False,
        limit: int = 100,
    ) -> list[dict[str, Any]]:
        stmt = select(SolutionRecordModel).order_by(SolutionRecordModel.updated_at.desc(), SolutionRecordModel.id.desc())
        if module:
            module_ids = self._find_solution_ids_by_module(module)
            stmt = stmt.where(or_(SolutionRecordModel.module.contains(module), SolutionRecordModel.id.in_(module_ids or [-1])))
        if submodule:
            stmt = stmt.where(SolutionRecordModel.submodule.contains(submodule))
        if error_code:
            stmt = stmt.where(SolutionRecordModel.error_code.contains(error_code))
        if error_name:
            stmt = stmt.where(SolutionRecordModel.error_name.contains(error_name))
        if message:
            stmt = stmt.where(SolutionRecordModel.message.contains(message))
        if message_keyword:
            keyword_ids = self._find_solution_ids_by_keyword(message_keyword)
            stmt = stmt.where(SolutionRecordModel.id.in_(keyword_ids or [-1]))
        if normalized_signature:
            stmt = stmt.where(SolutionRecordModel.normalized_signature == normalized_signature)
        if trigger_scenario:
            stmt = stmt.where(SolutionRecordModel.trigger_scenario.contains(trigger_scenario))
        if task_cluster:
            cluster_ids = self._find_solution_ids_by_task_cluster(task_cluster)
            stmt = stmt.where(SolutionRecordModel.id.in_(cluster_ids or [-1]))
        if submitter:
            stmt = stmt.where(SolutionRecordModel.submitter.contains(submitter))
        if reusable is not None:
            stmt = stmt.where(SolutionRecordModel.reusable == bool(reusable))
        if review_status:
            stmt = stmt.where(SolutionRecordModel.review_status == review_status)
        elif not viewer_is_reviewer:
            if viewer_username:
                stmt = stmt.where(or_(SolutionRecordModel.review_status == "approved", SolutionRecordModel.submitter == viewer_username))
            else:
                stmt = stmt.where(SolutionRecordModel.review_status == "approved")

        created_from_dt = self._parse_datetime(created_from)
        created_to_dt = self._parse_datetime(created_to)
        updated_from_dt = self._parse_datetime(updated_from)
        updated_to_dt = self._parse_datetime(updated_to)
        if created_from_dt:
            stmt = stmt.where(SolutionRecordModel.created_at >= created_from_dt)
        if created_to_dt:
            stmt = stmt.where(SolutionRecordModel.created_at <= created_to_dt)
        if updated_from_dt:
            stmt = stmt.where(SolutionRecordModel.updated_at >= updated_from_dt)
        if updated_to_dt:
            stmt = stmt.where(SolutionRecordModel.updated_at <= updated_to_dt)

        if search:
            fts_ids = self._search_solution_ids(search)
            if fts_ids:
                stmt = stmt.where(SolutionRecordModel.id.in_(fts_ids))
            else:
                stmt = stmt.where(
                    or_(
                        SolutionRecordModel.error_name.contains(search),
                        SolutionRecordModel.message.contains(search),
                        SolutionRecordModel.normalized_signature.contains(search),
                        SolutionRecordModel.error_code.contains(search),
                        SolutionRecordModel.trigger_scenario.contains(search),
                        SolutionRecordModel.root_cause_analysis.contains(search),
                        SolutionRecordModel.verified_solution.contains(search),
                        SolutionRecordModel.submitter.contains(search),
                    )
                )

        rows = list(self.db.scalars(stmt.limit(max(1, min(limit, 500)))))
        relation_maps = self._load_relation_maps([row.id for row in rows])
        return [self.serialize(row, relation_maps=relation_maps) for row in rows]

    def serialize(self, row: SolutionRecordModel, *, relation_maps: dict[str, dict[int, Any]] | None = None) -> dict[str, Any]:
        relation_maps = relation_maps or self._load_relation_maps([row.id])
        return {
            "id": row.id,
            "error_name": row.error_name,
            "error_category": row.error_category,
            "module": row.module,
            "submodule": row.submodule,
            "error_code": row.error_code,
            "error_code_prefix": row.error_code_prefix,
            "message": row.message,
            "normalized_signature": row.normalized_signature,
            "exception_description": row.exception_description,
            "trigger_scenario": row.trigger_scenario,
            "impact_scope": row.impact_scope,
            "report_source": row.report_source,
            "related_logs": _json_load(row.related_logs),
            "related_source_files": _json_load(row.related_source_files),
            "root_cause_analysis": row.root_cause_analysis,
            "verified_solution": row.verified_solution,
            "workaround": row.workaround,
            "owner_department": row.owner_department,
            "submitter": row.submitter,
            "source": row.source,
            "review_status": row.review_status,
            "reusable": row.reusable,
            "similar_case_refs": _json_load(row.similar_case_refs),
            "metadata": _json_load(row.metadata_json),
            "tags": relation_maps.get("tags", {}).get(row.id, []),
            "task_clusters": relation_maps.get("clusters", {}).get(row.id, []),
            "message_keywords": relation_maps.get("keywords", {}).get(row.id, []),
            "task_links": relation_maps.get("task_links", {}).get(row.id, []),
            "module_links": relation_maps.get("modules", {}).get(row.id, []),
            "created_at": row.created_at.isoformat() if row.created_at else None,
            "updated_at": row.updated_at.isoformat() if row.updated_at else None,
            "trigger_scenario_summary": _preview_text(row.trigger_scenario),
            "root_cause_summary": _preview_text(row.root_cause_analysis),
            "solution_summary": _preview_text(row.verified_solution),
        }

    def export_records(self, export_format: str = "json") -> str:
        rows = self.list_records(limit=5000, viewer_is_reviewer=True)
        export_dir = Path(self.settings.export_dir)
        export_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.utcnow().strftime("%Y%m%d%H%M%S")
        export_format = export_format.lower()

        if export_format == "json":
            path = export_dir / f"solution_repository_{timestamp}.json"
            path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
            return str(path)

        if export_format == "csv":
            path = export_dir / f"solution_repository_{timestamp}.csv"
            self._write_csv(path, rows)
            return str(path)

        if export_format in {"xlsx", "excel"}:
            from openpyxl import Workbook

            path = export_dir / f"solution_repository_{timestamp}.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "solutions"
            self._write_sheet(worksheet, rows)
            workbook.save(path)
            return str(path)

        if export_format == "sqlite":
            path = export_dir / f"solution_repository_{timestamp}.sqlite"
            self._write_sqlite_backup(path, rows)
            return str(path)

        raise ValueError(f"unsupported export format: {export_format}")

    def _sync_relations(self, solution_id: int, *, data: dict[str, Any], actor: str) -> None:
        self.db.query(SolutionModuleLinkModel).filter(SolutionModuleLinkModel.solution_id == solution_id).delete()
        self.db.query(SolutionTaskClusterLinkModel).filter(SolutionTaskClusterLinkModel.solution_id == solution_id).delete()
        self.db.query(SolutionTagLinkModel).filter(SolutionTagLinkModel.solution_id == solution_id).delete()
        self.db.query(SolutionMessageKeywordModel).filter(SolutionMessageKeywordModel.solution_id == solution_id).delete()
        self.db.query(SolutionTaskLinkModel).filter(SolutionTaskLinkModel.solution_id == solution_id).delete()

        primary_module = self.catalog_service.resolve_module(data["module"])
        module_names = [primary_module.module_key] + [name for name in data.get("additional_modules", []) if name != primary_module.module_key]
        for index, module_name in enumerate(module_names):
            try:
                module_row = self.catalog_service.resolve_module(module_name)
            except Exception:
                continue
            self.db.add(
                SolutionModuleLinkModel(
                    solution_id=solution_id,
                    module_key=module_row.module_key,
                    module_display_name=module_row.display_name,
                    module_config_id=module_row.id,
                    is_primary=index == 0,
                    created_at=utcnow(),
                )
            )

        for index, cluster_name in enumerate(data.get("task_clusters", [])):
            cluster_row = self.db.scalar(
                select(TaskClusterModel).where(
                    or_(TaskClusterModel.cluster_key == cluster_name, TaskClusterModel.display_name == cluster_name)
                )
            )
            if not cluster_row:
                cluster_payload = self.catalog_service.create_task_cluster(
                    display_name=cluster_name,
                    description=None,
                    actor=actor,
                    auto_approve=True,
                )
                cluster_row = self.db.get(TaskClusterModel, cluster_payload["id"])
            self.db.add(
                SolutionTaskClusterLinkModel(
                    solution_id=solution_id,
                    cluster_id=cluster_row.id,
                    is_primary=index == 0,
                    created_at=utcnow(),
                )
            )

        for tag_name in data.get("tags", []):
            tag_row = self.db.scalar(select(SolutionTagModel).where(SolutionTagModel.name == tag_name))
            if not tag_row:
                tag_row = SolutionTagModel(name=tag_name, created_at=utcnow())
                self.db.add(tag_row)
                self.db.flush()
            self.db.add(SolutionTagLinkModel(solution_id=solution_id, tag_id=tag_row.id, created_at=utcnow()))

        for keyword in data.get("message_keywords", []):
            if keyword:
                self.db.add(SolutionMessageKeywordModel(solution_id=solution_id, keyword=keyword.lower(), created_at=utcnow()))

        for link in data.get("task_links", []):
            if not isinstance(link, dict):
                continue
            task_uuid = str(link.get("task_uuid") or "").strip() or None
            normalized_signature = str(link.get("normalized_signature") or "").strip() or None
            if not task_uuid and not normalized_signature:
                continue
            self.db.add(
                SolutionTaskLinkModel(
                    solution_id=solution_id,
                    task_uuid=task_uuid,
                    normalized_signature=normalized_signature,
                    relation_type=str(link.get("relation_type") or "reference"),
                    source_note=str(link.get("source_note") or "").strip() or None,
                    created_by=actor,
                    created_at=utcnow(),
                )
            )

    def _sync_search_index(self, solution_id: int) -> None:
        if not self.fts_enabled():
            return
        payload = self.get_record(solution_id)
        document = " ".join(
            [
                str(payload.get("error_name") or ""),
                str(payload.get("error_code") or ""),
                str(payload.get("module") or ""),
                " ".join(payload.get("message_keywords") or []),
                " ".join(payload.get("tags") or []),
                " ".join(payload.get("task_clusters") or []),
                str(payload.get("message") or ""),
                str(payload.get("trigger_scenario") or ""),
                str(payload.get("root_cause_analysis") or ""),
                str(payload.get("verified_solution") or ""),
                str(payload.get("submitter") or ""),
                str(payload.get("review_status") or ""),
            ]
        )
        self.db.execute(text("DELETE FROM solution_search_fts WHERE solution_id = :solution_id"), {"solution_id": solution_id})
        self.db.execute(
            text("INSERT INTO solution_search_fts(solution_id, search_text) VALUES (:solution_id, :search_text)"),
            {"solution_id": solution_id, "search_text": document},
        )

    def _search_solution_ids(self, search_text: str) -> list[int]:
        if not self.fts_enabled():
            return []
        query = self._fts_query(search_text)
        if not query:
            return []
        try:
            rows = self.db.execute(
                text("SELECT solution_id FROM solution_search_fts WHERE solution_search_fts MATCH :query LIMIT 500"),
                {"query": query},
            ).fetchall()
            return [int(row[0]) for row in rows]
        except Exception:
            return []

    def _find_solution_ids_by_module(self, module_value: str) -> list[int]:
        rows = self.db.execute(
            select(SolutionModuleLinkModel.solution_id).where(
                or_(
                    SolutionModuleLinkModel.module_key.contains(module_value),
                    SolutionModuleLinkModel.module_display_name.contains(module_value),
                )
            )
        ).fetchall()
        return [int(row[0]) for row in rows]

    def _find_solution_ids_by_task_cluster(self, cluster_value: str) -> list[int]:
        rows = self.db.execute(
            select(SolutionTaskClusterLinkModel.solution_id)
            .join(TaskClusterModel, TaskClusterModel.id == SolutionTaskClusterLinkModel.cluster_id)
            .where(or_(TaskClusterModel.display_name.contains(cluster_value), TaskClusterModel.cluster_key.contains(cluster_value)))
        ).fetchall()
        return [int(row[0]) for row in rows]

    def _find_solution_ids_by_keyword(self, keyword_value: str) -> list[int]:
        rows = self.db.execute(
            select(SolutionMessageKeywordModel.solution_id).where(SolutionMessageKeywordModel.keyword.contains(keyword_value.lower()))
        ).fetchall()
        return [int(row[0]) for row in rows]

    def _load_relation_maps(self, solution_ids: list[int]) -> dict[str, dict[int, Any]]:
        if not solution_ids:
            return {"tags": {}, "clusters": {}, "keywords": {}, "task_links": {}, "modules": {}}
        tags_map: dict[int, list[str]] = {}
        cluster_map: dict[int, list[str]] = {}
        keyword_map: dict[int, list[str]] = {}
        task_link_map: dict[int, list[dict[str, Any]]] = {}
        module_map: dict[int, list[dict[str, Any]]] = {}

        tag_rows = self.db.execute(
            select(SolutionTagLinkModel.solution_id, SolutionTagModel.name)
            .join(SolutionTagModel, SolutionTagModel.id == SolutionTagLinkModel.tag_id)
            .where(SolutionTagLinkModel.solution_id.in_(solution_ids))
        ).fetchall()
        for solution_id, tag_name in tag_rows:
            tags_map.setdefault(int(solution_id), []).append(str(tag_name))

        cluster_rows = self.db.execute(
            select(SolutionTaskClusterLinkModel.solution_id, TaskClusterModel.display_name)
            .join(TaskClusterModel, TaskClusterModel.id == SolutionTaskClusterLinkModel.cluster_id)
            .where(SolutionTaskClusterLinkModel.solution_id.in_(solution_ids))
        ).fetchall()
        for solution_id, cluster_name in cluster_rows:
            cluster_map.setdefault(int(solution_id), []).append(str(cluster_name))

        keyword_rows = self.db.execute(
            select(SolutionMessageKeywordModel.solution_id, SolutionMessageKeywordModel.keyword).where(
                SolutionMessageKeywordModel.solution_id.in_(solution_ids)
            )
        ).fetchall()
        for solution_id, keyword in keyword_rows:
            keyword_map.setdefault(int(solution_id), []).append(str(keyword))

        task_link_rows = self.db.execute(
            select(
                SolutionTaskLinkModel.solution_id,
                SolutionTaskLinkModel.task_uuid,
                SolutionTaskLinkModel.normalized_signature,
                SolutionTaskLinkModel.relation_type,
                SolutionTaskLinkModel.source_note,
            ).where(SolutionTaskLinkModel.solution_id.in_(solution_ids))
        ).fetchall()
        for solution_id, task_uuid, signature, relation_type, source_note in task_link_rows:
            task_link_map.setdefault(int(solution_id), []).append(
                {
                    "task_uuid": task_uuid,
                    "normalized_signature": signature,
                    "relation_type": relation_type,
                    "source_note": source_note,
                }
            )

        module_rows = self.db.execute(
            select(
                SolutionModuleLinkModel.solution_id,
                SolutionModuleLinkModel.module_key,
                SolutionModuleLinkModel.module_display_name,
                SolutionModuleLinkModel.is_primary,
            ).where(SolutionModuleLinkModel.solution_id.in_(solution_ids))
        ).fetchall()
        for solution_id, module_key, display_name, is_primary in module_rows:
            module_map.setdefault(int(solution_id), []).append(
                {
                    "module_key": module_key,
                    "display_name": display_name,
                    "is_primary": bool(is_primary),
                }
            )

        return {
            "tags": tags_map,
            "clusters": cluster_map,
            "keywords": keyword_map,
            "task_links": task_link_map,
            "modules": module_map,
        }

    @staticmethod
    def _parse_datetime(value: str | None) -> datetime | None:
        text_value = str(value or "").strip()
        if not text_value:
            return None
        try:
            return datetime.fromisoformat(text_value)
        except Exception:
            return None

    @staticmethod
    def _fts_query(search_text: str) -> str:
        tokens = _tokenize_keywords(search_text)
        if not tokens:
            return ""
        return " OR ".join(tokens[:8])

    @staticmethod
    def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
        if not rows:
            path.write_text("", encoding="utf-8")
            return
        with path.open("w", encoding="utf-8-sig", newline="") as file_obj:
            writer = csv.DictWriter(file_obj, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            for row in rows:
                writer.writerow({key: SolutionRepositoryService._export_value(value) for key, value in row.items()})

    @staticmethod
    def _write_sheet(worksheet, rows: list[dict[str, Any]]) -> None:
        if not rows:
            worksheet.append(["empty"])
            return
        headers = list(rows[0].keys())
        worksheet.append(headers)
        for row in rows:
            worksheet.append([SolutionRepositoryService._export_value(row.get(header)) for header in headers])

    @staticmethod
    def _write_sqlite_backup(path: Path, rows: list[dict[str, Any]]) -> None:
        if path.exists():
            path.unlink()
        connection = sqlite3.connect(path)
        try:
            connection.execute(
                """
                CREATE TABLE solution_records (
                    id INTEGER,
                    error_name TEXT,
                    error_category TEXT,
                    module TEXT,
                    submodule TEXT,
                    error_code TEXT,
                    error_code_prefix TEXT,
                    message TEXT,
                    normalized_signature TEXT,
                    exception_description TEXT,
                    trigger_scenario TEXT,
                    impact_scope TEXT,
                    report_source TEXT,
                    related_logs TEXT,
                    related_source_files TEXT,
                    root_cause_analysis TEXT,
                    verified_solution TEXT,
                    workaround TEXT,
                    owner_department TEXT,
                    submitter TEXT,
                    source TEXT,
                    review_status TEXT,
                    reusable INTEGER,
                    similar_case_refs TEXT,
                    metadata TEXT,
                    tags TEXT,
                    task_clusters TEXT,
                    message_keywords TEXT,
                    task_links TEXT,
                    module_links TEXT,
                    created_at TEXT,
                    updated_at TEXT
                )
                """
            )
            if rows:
                columns = list(rows[0].keys())
                placeholders = ",".join("?" for _ in columns)
                connection.executemany(
                    f"INSERT INTO solution_records ({','.join(columns)}) VALUES ({placeholders})",
                    [[SolutionRepositoryService._export_value(row.get(column)) for column in columns] for row in rows],
                )
            connection.commit()
        finally:
            connection.close()

    @staticmethod
    def _export_value(value: Any) -> Any:
        if value is None:
            return ""
        if isinstance(value, (str, int, float, bool)):
            return value
        return json.dumps(value, ensure_ascii=False, default=str)
